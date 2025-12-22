"""
PPTX presentation extraction
"""
from pptx import Presentation
from typing import Dict, List
import logging

logger = logging.getLogger(__name__)


class PPTXExtractor:
    """Extract content from PPTX files"""
    
    def extract(self, file_path: str) -> Dict:
        """
        Extract text, slide titles, and images from PPTX
        
        Args:
            file_path: Path to PPTX file
            
        Returns:
            Dictionary containing extracted content
        """
        try:
            prs = Presentation(file_path)
            
            result = {
                'text': '',
                'headings': [],
                'sections': [],
                'images': [],
                'slides': len(prs.slides),
                'metadata': {}
            }
            
            # Extract core properties
            if prs.core_properties:
                result['metadata'] = {
                    'title': prs.core_properties.title or '',
                    'author': prs.core_properties.author or '',
                    'subject': prs.core_properties.subject or ''
                }
            
            # Extract content from each slide
            full_text = []
            
            for slide_idx, slide in enumerate(prs.slides, 1):
                slide_text = []
                slide_title = None
                
                # Extract text from shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text = shape.text.strip()
                        
                        # First text box is usually the title
                        if slide_title is None and text:
                            slide_title = text
                            result['headings'].append({
                                'text': text,
                                'level': 2,  # Slide titles as H2
                                'slide': slide_idx
                            })
                        else:
                            slide_text.append(text)
                    
                    # Check for images
                    if hasattr(shape, "image"):
                        result['images'].append({
                            'slide': slide_idx,
                            'type': 'image'
                        })
                
                # Create section for this slide
                if slide_title:
                    result['sections'].append({
                        'title': slide_title,
                        'level': 2,
                        'slide': slide_idx,
                        'content': '\n'.join(slide_text)
                    })
                    full_text.append(f"# {slide_title}\n" + '\n'.join(slide_text))
                else:
                    full_text.append('\n'.join(slide_text))
            
            result['text'] = '\n\n'.join(full_text)
            
            logger.info(f"Extracted {len(result['headings'])} slide titles from {result['slides']} slides")
            
            return result
            
        except Exception as e:
            logger.error(f"Error extracting PPTX: {e}")
            raise


def extract_pptx(file_path: str) -> Dict:
    """
    Convenience function to extract PPTX content
    
    Args:
        file_path: Path to PPTX file
        
    Returns:
        Extracted content dictionary
    """
    extractor = PPTXExtractor()
    return extractor.extract(file_path)
